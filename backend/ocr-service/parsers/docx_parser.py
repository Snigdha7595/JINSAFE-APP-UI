from docx import Document
import logging
from typing import Dict, Any
import io

logger = logging.getLogger(__name__)

class DocxParser:
    """Parser for DOCX files"""
    
    async def parse(self, file_content: bytes) -> Dict[str, Any]:
        """Parse DOCX file and extract text"""
        try:
            doc = Document(io.BytesIO(file_content))
            
            extracted_text = ""
            for para in doc.paragraphs:
                extracted_text += para.text + "\n"
            
            return {
                "text": extracted_text,
                "page_count": 1,
                "format": "docx",
                "paragraph_count": len(doc.paragraphs)
            }
        except Exception as e:
            logger.error(f"DOCX parsing failed: {str(e)}")
            raise
